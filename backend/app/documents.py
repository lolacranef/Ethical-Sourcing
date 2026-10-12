import csv
import io
from pathlib import PurePath

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader


class UnsupportedDocumentError(ValueError):
    pass


def extract_document_text(filename: str, content: bytes) -> tuple[str, list[dict]]:
    extension = PurePath(filename).suffix.casefold()
    if extension in {".txt", ".md"}:
        text = content.decode("utf-8-sig", errors="replace")
        return text, [{"page": 1, "start": 0, "end": len(text)}]
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        pages = []
        offset = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append({"page": page_number, "text": text, "start": offset, "end": offset + len(text)})
            offset += len(text) + 1
        full_text = "\n".join(page["text"] for page in pages)
        if not full_text.strip():
            raise UnsupportedDocumentError(
                "This PDF has no extractable text. OCR for scanned PDFs is not enabled in this MVP."
            )
        return full_text, pages
    if extension == ".docx":
        doc = WordDocument(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        return text, [{"page": None, "start": 0, "end": len(text)}]
    if extension in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = [" | ".join(str(value) for value in row if value is not None) for sheet in workbook for row in sheet.iter_rows(values_only=True)]
        text = "\n".join(row for row in rows if row)
        return text, [{"page": None, "start": 0, "end": len(text)}]
    if extension == ".csv":
        rows = csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace")))
        text = "\n".join(" | ".join(row) for row in rows)
        return text, [{"page": None, "start": 0, "end": len(text)}]
    raise UnsupportedDocumentError(f"Unsupported document type: {extension or 'no extension'}")


def classify_document(text: str) -> str:
    lowered = text.casefold()
    if any(term in lowered for term in ("sanctions notice", "designated pursuant", "entity list")):
        return "sanctions_notice"
    if any(term in lowered for term in ("court", "plaintiff", "defendant", "case no.")):
        return "court_filing"
    if any(term in lowered for term in ("customs declaration", "bill of lading", "importer of record")):
        return "customs_record"
    if any(term in lowered for term in ("know your customer", "customer due diligence", "kyc")):
        return "kyc_file"
    if any(term in lowered for term in ("ngo report", "labor rights", "human rights report")):
        return "ngo_report"
    return "other"